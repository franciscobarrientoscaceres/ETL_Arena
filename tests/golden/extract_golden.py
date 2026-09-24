#!/usr/bin/env python3
"""
extract_golden.py
=================
Script para extraer un golden reference desde el Excel y guardarlo en
tests/golden/data/golden_YYYY_MM_<name>.json

Uso:
    python tests/golden/extract_golden.py \
        --excel "data/AvailabilityCalculation_PCS&Batteries_20260907_septiembre 2026.xlsm" \
        --month 8 --year 2026 --label "august"

Prerequisito: el libro debe tener la macro VBA ya ejecutada para el mes
indicado (C5=periodo_inicio, C7=periodo_fin, C12 y C14 ya calculados).
"""
import argparse
import json
import pathlib
from datetime import datetime
import openpyxl


def extract(excel_path: str, year: int, month: int, label: str) -> dict:
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    ws_calc   = wb["Calculation-Availability"]
    ws_daily  = wb["Daily"]
    ws_faults = wb["ListOfFaults"]
    ws_annual = wb["Annual_AVA"]

    def dt(v):
        return v.strftime("%Y-%m-%d %H:%M:%S") if isinstance(v, datetime) else str(v) if v else None

    # KPI del periodo
    period = {
        "year": year,
        "month": month,
        "month_name": label,
        "period_start": dt(ws_calc.cell(5, 3).value)[:10] if ws_calc.cell(5, 3).value else None,
        "period_end":   dt(ws_calc.cell(7, 3).value)[:10] if ws_calc.cell(7, 3).value else None,
        "parameters": {
            "total_pcs":             int(ws_calc.cell(2,  3).value),
            "batteries_per_pcs":     int(ws_calc.cell(3,  3).value),
            "total_racks":           int(ws_calc.cell(11, 3).value),
            "sampling_minutes":      int(ws_calc.cell(23, 3).value),
            "only_operational_time": ws_calc.cell(21, 3).value,
            "apply_excused_event":   ws_calc.cell(31, 3).value,
            "algorithm_version":     "availability-v1-excel-parity",
            "source_file":           pathlib.Path(excel_path).name,
            "extraction_date":       datetime.now().strftime("%Y-%m-%d"),
        },
        "kpi": {
            "c12_sample_blocks":           int(ws_calc.cell(12, 3).value),
            "c14_unavailable_rack_blocks": ws_calc.cell(14, 3).value,
            "c16_availability_period":     ws_calc.cell(16, 3).value,
            "c19_accumulated_annual":      ws_calc.cell(19, 3).value,
        },
    }

    # Daily
    daily = []
    for row in range(9, 40):
        c = ws_daily.cell(row, 3).value
        if c is None:
            break
        daily.append({
            "day_number": row - 8,
            "date": dt(c)[:10] if c else None,
            "daily_unavailable_rack_blocks":        ws_daily.cell(row, 4).value,
            "accumulated_unavailable_rack_blocks":  ws_daily.cell(row, 5).value,
            "availability":                         ws_daily.cell(row, 6).value,
            "variation":                            ws_daily.cell(row, 7).value,
        })

    # Annual_AVA snapshot
    annual = {}
    for row in range(7, 19):
        mn = ws_annual.cell(row, 2).value
        if not mn:
            continue
        unav  = ws_annual.cell(row, 6).value
        avail = ws_annual.cell(row, 7).value
        if unav is None and avail is None:
            continue
        annual[str(int(mn))] = {
            "month_name":                     ws_annual.cell(row, 3).value,
            "days_in_month":                  ws_annual.cell(row, 4).value,
            "sampling_blocks":                ws_annual.cell(row, 5).value,
            "unavailable_rack_blocks":        unav,
            "monthly_availability":           avail,
            "accumulated_sampling_blocks":    ws_annual.cell(row, 8).value,
            "accumulated_unavailable_blocks": ws_annual.cell(row, 9).value,
        }

    # Fault events
    total_rh = ws_faults.cell(10, 12).value
    events = []
    for r in ws_faults.iter_rows(min_row=6, max_row=300, values_only=True):
        b, c, d, e, f, g, h, i = r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8]
        if b is None:
            continue
        events.append({
            "pcs_number":                 int(b),
            "start_timestamp":            dt(c),
            "end_timestamp":              dt(d),
            "duration_hours":             e,
            "fault_description":          f,
            "fault_code":                 g,
            "average_batteries_involved": h,
            "unavailable_rack_hours":     i,
        })

    return {
        "_meta": {
            "description":          "Golden reference extraida directamente del Excel con openpyxl data_only=True",
            "extraction_method":    "openpyxl read_only=True, data_only=True",
            "extraction_date":      datetime.now().strftime("%Y-%m-%d"),
            "tolerance_internal":   1e-9,
            "tolerance_kpi":        1e-6,
            "tolerance_rack_hours": 1e-3,
            "note":                 f"Corrida VBA con Only Operational Time={period['parameters']['only_operational_time']}, Excusable Event={period['parameters']['apply_excused_event']}",
        },
        "period": period,
        "daily":  daily,
        "annual_ava_snapshot": annual,
        "fault_events_summary": {
            "total_events":                   len(events),
            "total_unavailable_rack_hours":   total_rh,
            "events_first_20":                events[:20],
        },
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", required=True)
    ap.add_argument("--month", required=True, type=int)
    ap.add_argument("--year",  required=True, type=int)
    ap.add_argument("--label", required=True, help="Nombre del mes en minusculas, ej: july")
    args = ap.parse_args()

    data = extract(args.excel, args.year, args.month, args.label)

    out_dir = pathlib.Path(__file__).parent / "data"
    out_dir.mkdir(exist_ok=True)
    out_file = out_dir / f"golden_{args.year}_{args.month:02d}_{args.label}.json"
    out_file.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Guardado: {out_file}")
    kpi = data["period"]["kpi"]
    print(f"  C12={kpi['c12_sample_blocks']}  C16={kpi['c16_availability_period']:.8f}")
    print(f"  Events={data['fault_events_summary']['total_events']}  RackHours={data['fault_events_summary']['total_unavailable_rack_hours']:.4f}")


if __name__ == "__main__":
    main()
